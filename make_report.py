from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

GREEN_DARK  = RGBColor(0x1a, 0x3a, 0x2a)
GREEN_MID   = RGBColor(0x2d, 0x6a, 0x4f)
GREEN_LIGHT = RGBColor(0xd8, 0xf3, 0xdc)
AMBER       = RGBColor(0xf5, 0x9e, 0x0b)
AMBER_LIGHT = RGBColor(0xff, 0xfb, 0xeb)
RED         = RGBColor(0xc0, 0x39, 0x2b)
GRAY        = RGBColor(0x55, 0x55, 0x55)
WHITE       = RGBColor(0xff, 0xff, 0xff)
YELLOW_BG   = RGBColor(0xff, 0xf9, 0xc3)

def set_cell_bg(cell, rgb: RGBColor):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    hex_color = f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
    shd.set(qn('w:fill'), hex_color)
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)

def heading(text, level=1, color=GREEN_DARK):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16 if level == 1 else 13)
    run.font.color.rgb = color
    return p

def para(text, bold=False, color=None, size=10, italic=False, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p

def section_label(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = GREEN_MID

# ── TITLE ──────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
run = p.add_run('Test kreacji — „O zwierzętach, które wybrały Tatry”')
run.bold = True
run.font.size = Pt(20)
run.font.color.rgb = GREEN_DARK

para("Tatrzański Park Narodowy  ·  Meta Ads  ·  Test AB  ·  20–27 maja 2026",
     color=GRAY, size=10)

# Meta row
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(12)
for label, val in [("Budżet:", "~797 PLN"), ("Cel:", "Sprzedaż"),
                   ("Odbiorcy:", "Zainteresowania"), ("Warianty:", "3 kreacje")]:
    p.add_run(f"{label} ").bold = False
    r = p.add_run(f"{val}    ")
    r.bold = True
    r.font.size = Pt(10)

# ── SCALE NOTE ─────────────────────────────────────────────────────────────
section_label("Uwaga metodologiczna")
t = doc.add_table(rows=1, cols=1)
t.style = 'Table Grid'
cell = t.rows[0].cells[0]
set_cell_bg(cell, YELLOW_BG)
cp = cell.paragraphs[0]
r = cp.add_run("⚠️  Dlaczego skupiamy się na CPA, nie ROAS?\n")
r.bold = True
r.font.size = Pt(10)
r2 = cp.add_run(
    "Przy 6–13 zakupach na wariant jeden droższy zakup może całkowicie wywrócić ROAS "
    "i średnią wartość zamówienia. Na tej skali są to dane przypadkowe, nie trendy. "
    "CPA (koszt pozyskania zakupu) jest tu dużo bardziej wiarygodną miarą."
)
r2.font.size = Pt(10)
doc.add_paragraph()

# ── PART 1 ─────────────────────────────────────────────────────────────────
section_label("Część 1 z 2")
heading("Plener vs Studio  (równy budżet ~400 PLN)", level=1)

headers = ["Metryka", "🌿 Plener", "🎬 Studio łącznie", "Lepszy wynik"]
rows = [
    ("Wydano (PLN)",     "398,47",  "398,98",  "—"),
    ("Wyświetlenia",     "54 115",  "46 278",  "🌿 Plener"),
    ("CPM (PLN)",        "7,36",    "8,62",    "🌿 Plener"),
    ("CTR link",         "1,09%",   "0,76%",   "🌿 Plener"),
    ("CPC link (PLN)",   "0,68",    "1,13",    "🌿 Plener"),
    ("Zakupy",           "13",      "13",      "— Remis"),
    ("⭐ CPA (PLN)",     "30,65",   "30,69",   "— Remis"),
]

t1 = doc.add_table(rows=len(rows)+1, cols=4)
t1.style = 'Table Grid'

# header row
for i, h in enumerate(headers):
    cell = t1.rows[0].cells[i]
    set_cell_bg(cell, GREEN_DARK)
    p2 = cell.paragraphs[0]
    r = p2.add_run(h)
    r.bold = True
    r.font.color.rgb = WHITE
    r.font.size = Pt(10)

for ri, row in enumerate(rows):
    is_cpa = row[0].startswith("⭐")
    for ci, val in enumerate(row):
        cell = t1.rows[ri+1].cells[ci]
        if is_cpa:
            set_cell_bg(cell, GREEN_LIGHT)
        p2 = cell.paragraphs[0]
        r = p2.add_run(val)
        r.font.size = Pt(10)
        if is_cpa:
            r.bold = True
        # colour best/worst in data cols
        if ci == 1 and row[3] == "🌿 Plener":
            r.font.color.rgb = GREEN_MID
            r.bold = True
        if ci == 2 and row[3] == "🎬 Studio":
            r.font.color.rgb = GREEN_MID
            r.bold = True

doc.add_paragraph()
para("Uwaga: Plener otrzymał 2× wyższy budżet niż każdy wariant studyjny z osobna. "
     "Studio = suma obu wariantów — budżety wyrównane do ~400 PLN.",
     italic=True, color=GRAY, size=9)

# ── PART 2 ─────────────────────────────────────────────────────────────────
section_label("Część 2 z 2")
heading("Studio: białe tło vs czarne tło", level=1)

headers2 = ["Metryka", "⬜ Białe tło", "⬛ Czarne tło", "Lepszy wynik"]
rows2 = [
    ("Wydano (PLN)",    "199,72",  "199,26",  "—"),
    ("CPM (PLN)",       "9,47",    "7,91",    "⬛ Czarne"),
    ("CTR link",        "0,88%",   "0,67%",   "⬜ Białe"),
    ("CPC link (PLN)",  "1,08",    "1,19",    "⬜ Białe"),
    ("Zakupy",          "6",       "7",       "⬛ Czarne"),
    ("⭐ CPA (PLN)",    "33,29",   "28,47",   "⬛ Czarne"),
]

t2 = doc.add_table(rows=len(rows2)+1, cols=4)
t2.style = 'Table Grid'

for i, h in enumerate(headers2):
    cell = t2.rows[0].cells[i]
    set_cell_bg(cell, GREEN_DARK)
    p2 = cell.paragraphs[0]
    r = p2.add_run(h)
    r.bold = True
    r.font.color.rgb = WHITE
    r.font.size = Pt(10)

for ri, row in enumerate(rows2):
    is_cpa = row[0].startswith("⭐")
    for ci, val in enumerate(row):
        cell = t2.rows[ri+1].cells[ci]
        if is_cpa:
            set_cell_bg(cell, GREEN_LIGHT)
        p2 = cell.paragraphs[0]
        r = p2.add_run(val)
        r.font.size = Pt(10)
        if is_cpa:
            r.bold = True
        if ci == 1 and row[3] == "⬜ Białe":
            r.font.color.rgb = GREEN_MID
            r.bold = True
        if ci == 2 and row[3] == "⬛ Czarne":
            r.font.color.rgb = GREEN_MID
            r.bold = True
        if is_cpa and ci == 1:
            r.font.color.rgb = RED
        if is_cpa and ci == 2:
            r.font.color.rgb = GREEN_MID

doc.add_paragraph()

# ── WINNERS ────────────────────────────────────────────────────────────────
section_label("Podsumowanie")
heading("Kto wygrał?", level=1)

win_table = doc.add_table(rows=1, cols=2)
win_table.style = 'Table Grid'

# Runda 1
c1 = win_table.rows[0].cells[0]
set_cell_bg(c1, GREEN_LIGHT)
p2 = c1.paragraphs[0]
r = p2.add_run("Runda 1 — Plener vs Studio\n")
r.font.size = Pt(9); r.font.color.rgb = GRAY
r2 = p2.add_run("🤝 Remis na CPA\n")
r2.bold = True; r2.font.size = Pt(13); r2.font.color.rgb = GREEN_DARK
r3 = p2.add_run(
    "Przy równym budżecie i tej samej liczbie zakupów (13 vs 13) CPA jest "
    "praktycznie identyczny: 30,65 vs 30,69 PLN. Plener ma przewagę w CTR i CPC — "
    "tańszy ruch, ale nie przekłada się to na więcej zakupów.\n\n"
)
r3.font.size = Pt(10)
r4 = p2.add_run("CPA Plener: 30,65 PLN   |   CPA Studio: 30,69 PLN")
r4.bold = True; r4.font.size = Pt(10); r4.font.color.rgb = GREEN_MID

# Runda 2
c2 = win_table.rows[0].cells[1]
set_cell_bg(c2, AMBER_LIGHT)
p2 = c2.paragraphs[0]
r = p2.add_run("Runda 2 — Białe vs Czarne tło\n")
r.font.size = Pt(9); r.font.color.rgb = GRAY
r2 = p2.add_run("⬛ Czarne tło wygrywa CPA\n")
r2.bold = True; r2.font.size = Pt(13); r2.font.color.rgb = RGBColor(0x92, 0x40, 0x0e)
r3 = p2.add_run(
    "Czarne tło pozyskało zakup o 4,82 PLN taniej (28,47 vs 33,29 PLN) i wygenerowało "
    "o jeden zakup więcej przy identycznym budżecie. Różnica widoczna, choć skala "
    "(6 vs 7 zakupów) wymaga ostrożności.\n\n"
)
r3.font.size = Pt(10)
r4 = p2.add_run("CPA czarne: 28,47 PLN   |   CPA białe: 33,29 PLN")
r4.bold = True; r4.font.size = Pt(10); r4.font.color.rgb = RGBColor(0x92, 0x40, 0x0e)

doc.add_paragraph()

# ── HIERARCHY ──────────────────────────────────────────────────────────────
section_label("Hierarchia wyników (według CPA)")
for rank, name, detail, tag in [
    ("🥇", "Studio — czarne tło",  "CPA 28,47 PLN",              "Testuj dalej / remarketing"),
    ("🥈", "Plener",               "CPA 30,65 PLN  ·  CTR 1,09%","Dobry do ruchu / TOF"),
    ("🥉", "Studio — białe tło",   "CPA 33,29 PLN",              "Wymagał wyższego CPA"),
]:
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    p2.add_run(f"{rank}  ").font.size = Pt(11)
    r = p2.add_run(name)
    r.bold = True; r.font.size = Pt(11)
    r2 = p2.add_run(f"   {detail}   ")
    r2.font.size = Pt(10); r2.font.color.rgb = GRAY
    r3 = p2.add_run(f"[{tag}]")
    r3.font.size = Pt(9); r3.font.color.rgb = GREEN_MID; r3.bold = True

doc.add_paragraph()

# ── RECOMMENDATIONS ────────────────────────────────────────────────────────
section_label("Następne kroki")
heading("Rekomendacje", level=1)

for icon, title, text in [
    ("🔁", "Powtórz test na większej skali",
     "Przy 6–13 zakupach wnioski są kierunkowe, nie pewne. Powtórz test "
     "z budżetem min. 600–800 PLN na wariant, by osiągnąć 20+ zakupów i "
     "statystyczną wiarygodność."),
    ("⬛", "Czarne tło — priorytet",
     "Najniższy CPA w teście. Pierwszy kandydat do kolejnego testu i do "
     "remarketingu na osoby, które odwiedziły sklep ale nie kupiły."),
    ("🌿", "Plener — cold traffic / TOF",
     "Najlepszy CTR i najtańszy ruch. Przy remisie na CPA warto go zostawić "
     "jako kreację do zimnych odbiorców, gdzie zasięg i koszt kliknięcia są ważne."),
]:
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(6)
    r = p2.add_run(f"{icon}  {title}\n")
    r.bold = True; r.font.size = Pt(11); r.font.color.rgb = GREEN_DARK
    r2 = p2.add_run(text)
    r2.font.size = Pt(10)

doc.add_paragraph()
para("Tatrzański Park Narodowy  ·  Meta Ads  ·  202605_O zwierzętach – test AB  ·  3 czerwca 2026",
     color=GRAY, size=9, italic=True)

doc.save("/home/user/claude-ads/raport-test-ab-zwierzeta.docx")
print("done")
